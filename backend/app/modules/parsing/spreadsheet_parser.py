"""XLSX / XLS / CSV parsing with bounded memory and exact cell provenance.

Policies:
* Formulas are never evaluated. XLSX formulas are reported together with the
  cached value written by the authoring application; a missing cached value is
  reported as an issue and the cell value is left empty.
* Merged cells: the value lives only in the top-left cell (openpyxl read-only
  semantics). Values are NOT propagated into the other merged cells, to avoid
  double counting amounts.
* Empty rows are skipped but row numbers always refer to the original sheet.
* Macros are rejected at upload; parsers never execute workbook code.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from openpyxl.utils import get_column_letter

from app.core.config import get_settings
from app.core.exceptions import PermanentProcessingError
from app.shared.enums import FileKind
from app.shared.text import fold


class SpreadsheetMalformedError(PermanentProcessingError):
    code = "SPREADSHEET_MALFORMED"


class SpreadsheetTooLargeError(PermanentProcessingError):
    code = "SPREADSHEET_TOO_LARGE"


FIELD_SYNONYMS: dict[str, tuple[str, ...]] = {
    "invoice_number": (
        "invoice no",
        "invoice number",
        "invoice #",
        "inv no",
        "invoice id",
        "invoice",
        "bill no",
        "document no",
        "doc no",
        "invoice ref",
    ),
    "supplier_name": ("supplier", "vendor", "vendor name", "supplier name", "seller", "payee", "creditor"),
    "customer_name": ("customer", "client", "buyer", "bill to", "customer name", "client name"),
    "invoice_date": ("invoice date", "date", "inv date", "document date", "bill date", "issue date"),
    "currency": ("currency", "curr", "ccy", "currency code"),
    "subtotal": ("subtotal", "sub total", "net amount", "net", "amount excl tax", "amount before tax", "net total"),
    "tax_amount": ("tax", "vat", "gst", "tax amount", "tax value", "vat amount", "sales tax", "tax total"),
    "total_amount": (
        "total",
        "invoice total",
        "gross amount",
        "total amount",
        "grand total",
        "amount due",
        "invoice amount",
        "gross total",
        "amount",
    ),
    "purchase_order_number": ("po", "po number", "purchase order", "po no"),
    "payment_due_date": ("due date", "payment due", "due"),
    "supplier_tax_id": ("vat id", "tax id", "supplier vat", "vat number", "tin"),
}

_HEADER_CLEAN_RE = re.compile(r"[^0-9a-z#]+")


def normalize_header(value: str) -> str:
    return _HEADER_CLEAN_RE.sub(" ", fold(value)).strip()


def header_field_score(header: str) -> tuple[str | None, float]:
    """Best invoice field for a header and a string-similarity score (0..1)."""
    from rapidfuzz import fuzz

    norm = normalize_header(header)
    if not norm:
        return None, 0.0
    best: tuple[str | None, float] = (None, 0.0)
    for field_name, synonyms in FIELD_SYNONYMS.items():
        if norm == field_name.replace("_", " "):
            return field_name, 1.0
        for syn in synonyms:
            score = 1.0 if norm == syn else fuzz.ratio(norm, syn) / 100
            if score > best[1]:
                best = (field_name, score)
    return best


@dataclass
class RawCell:
    value: Any
    formula: str | None = None
    missing_cached: bool = False


@dataclass
class SheetTable:
    sheet_name: str | None
    table_index: int
    header_row_number: int | None
    headers: list[str]
    issues: list[dict[str, Any]] = field(default_factory=list)
    column_count: int = 0
    truncated: bool = False
    looks_like_register: bool = False


def cell_to_json(value: Any) -> tuple[str | None, str]:
    """Normalize a cell value to (string, type). Floats use shortest repr."""
    if value is None:
        return None, "empty"
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), "bool"
    if isinstance(value, int):
        return str(value), "number"
    if isinstance(value, float):
        # Spreadsheet numbers are IEEE doubles at the source; repr() is the
        # shortest decimal string that round-trips, so no precision is invented.
        return format(Decimal(repr(value)), "f"), "number"
    if isinstance(value, Decimal):
        return format(value, "f"), "number"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat(), "date"
        return value.isoformat(), "datetime"
    if isinstance(value, date):
        return value.isoformat(), "date"
    if isinstance(value, time):
        return value.isoformat(), "time"
    text = str(value).strip()
    return (text or None), "text"


# --- Row sources -------------------------------------------------------------


def _iter_xlsx(path: str) -> Iterator[tuple[str, Iterator[tuple[int, list[RawCell]]]]]:
    from openpyxl import load_workbook

    try:
        wb_formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        wb_values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise SpreadsheetMalformedError(f"XLSX could not be opened: {type(exc).__name__}") from exc

    def rows(name: str) -> Iterator[tuple[int, list[RawCell]]]:
        ws_f, ws_v = wb_formulas[name], wb_values[name]
        start = ws_f.min_row or 1
        for offset, (frow, vrow) in enumerate(zip(ws_f.iter_rows(values_only=True), ws_v.iter_rows(values_only=True), strict=False)):
            cells: list[RawCell] = []
            for f_val, v_val in zip(frow, vrow, strict=False):
                if isinstance(f_val, str) and f_val.startswith("="):
                    cells.append(RawCell(value=v_val, formula=f_val[:500], missing_cached=v_val is None))
                else:
                    cells.append(RawCell(value=v_val))
            yield start + offset, cells

    try:
        for name in wb_formulas.sheetnames:
            yield name, rows(name)
    finally:
        wb_formulas.close()
        wb_values.close()


def _iter_xls(path: str) -> Iterator[tuple[str, Iterator[tuple[int, list[RawCell]]]]]:
    import xlrd

    try:
        book = xlrd.open_workbook(path, on_demand=True, formatting_info=False)
    except Exception as exc:
        raise SpreadsheetMalformedError(f"XLS could not be opened: {type(exc).__name__}") from exc

    def rows(sheet: Any) -> Iterator[tuple[int, list[RawCell]]]:
        for r in range(sheet.nrows):
            cells: list[RawCell] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(cell.value, book.datemode)
                    except Exception:  # noqa: BLE001
                        value = cell.value
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or cell.ctype == xlrd.XL_CELL_ERROR:
                    value = None
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                    value = int(value)
                cells.append(RawCell(value=value))
            yield r + 1, cells

    try:
        for idx in range(book.nsheets):
            sheet = book.sheet_by_index(idx)
            yield sheet.name, rows(sheet)
            book.unload_sheet(idx)
    finally:
        book.release_resources()


def detect_encoding(sample: bytes, candidates: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp1252", "latin-1")) -> str:
    for enc in candidates:
        try:
            sample.decode(enc)
            return enc
        except UnicodeDecodeError as exc:
            # Only tolerate a multi-byte character cut off at the end of the sample.
            if exc.reason == "unexpected end of data" and exc.start >= len(sample) - 4:
                return enc
            continue
    return "latin-1"


def _iter_csv(path: str, delimiter: str | None = None) -> Iterator[tuple[str, Iterator[tuple[int, list[RawCell]]]]]:
    with open(path, "rb") as fh:
        sample = fh.read(64 * 1024)
    encoding = detect_encoding(sample)
    text_sample = sample.decode(encoding, errors="replace")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text_sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","

    def rows() -> Iterator[tuple[int, list[RawCell]]]:
        with open(path, encoding=encoding, newline="", errors="replace") as fh:
            reader = csv.reader(fh, delimiter=delimiter, quotechar='"', doublequote=True, strict=False)
            try:
                for i, row in enumerate(reader, start=1):
                    yield i, [RawCell(value=v) for v in row]
            except csv.Error as exc:
                raise SpreadsheetMalformedError(f"CSV parse error near line {reader.line_num}") from exc

    yield "csv", rows()


def iter_sheets(path: str, kind: FileKind) -> Iterator[tuple[str, Iterator[tuple[int, list[RawCell]]]]]:
    if kind == FileKind.XLSX:
        return _iter_xlsx(path)
    if kind == FileKind.XLS:
        return _iter_xls(path)
    if kind == FileKind.CSV:
        return _iter_csv(path)
    raise ValueError(f"not a spreadsheet: {kind}")


# --- Table building ----------------------------------------------------------

HEADER_SCAN_ROWS = 30


def _header_score(cells: list[RawCell]) -> int:
    score = 0
    for cell in cells:
        if isinstance(cell.value, str) and cell.value.strip():
            _, s = header_field_score(cell.value)
            if s >= 0.85:
                score += 1
    return score


def _unique_headers(raw: list[Any], width: int) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        value = raw[i] if i < len(raw) else None
        name = str(value).strip() if value not in (None, "") else f"column_{i + 1}"
        name = name[:120]
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        headers.append(name)
    return headers


def build_tables(path: str, kind: FileKind) -> Iterator[tuple[SheetTable, Iterator[tuple[int, dict[str, Any]]]]]:
    """Yield (table, rows) per sheet. ``rows`` streams normalized data rows and
    must be fully consumed before advancing; ``table.issues`` and
    ``table.truncated`` are complete once ``rows`` is exhausted."""
    settings = get_settings()
    max_rows, max_cols = settings.max_spreadsheet_rows, settings.max_spreadsheet_columns
    counter = {"total": 0, "stop": False}
    for sheet_name, rows in iter_sheets(path, kind):
        if counter["stop"]:
            break
        buffered: list[tuple[int, list[RawCell]]] = []
        header_row: tuple[int, list[RawCell]] | None = None
        best_score = 0
        table = SheetTable(
            sheet_name=None if kind == FileKind.CSV else sheet_name,
            table_index=0,
            header_row_number=None,
            headers=[],
        )
        iterator = iter(rows)
        for row_number, cells in iterator:
            if len(cells) > max_cols:
                table.issues.append({"code": "TOO_MANY_COLUMNS", "row": row_number, "columns": len(cells)})
                cells = cells[:max_cols]
            buffered.append((row_number, cells))
            score = _header_score(cells)
            if score > best_score:
                best_score, header_row = score, (row_number, cells)
            if len(buffered) >= HEADER_SCAN_ROWS:
                break
        if header_row is None:
            header_row = next(
                ((n, c) for n, c in buffered if sum(1 for x in c if isinstance(x.value, str) and x.value.strip()) >= 2),
                None,
            )
        width = max((len(c) for _, c in buffered), default=0)
        if header_row is not None:
            table.header_row_number = header_row[0]
            table.headers = _unique_headers([c.value for c in header_row[1]], width)
        else:
            table.headers = _unique_headers([], width)
        table.column_count = width
        table.looks_like_register = best_score >= 2
        yield table, _stream_rows(table, kind, buffered, iterator, width, max_rows, max_cols, counter)


def _stream_rows(
    table: SheetTable,
    kind: FileKind,
    buffered: list[tuple[int, list[RawCell]]],
    iterator: Iterator[tuple[int, list[RawCell]]],
    width: int,
    max_rows: int,
    max_cols: int,
    counter: dict[str, Any],
) -> Iterator[tuple[int, dict[str, Any]]]:
    def source() -> Iterator[tuple[int, list[RawCell]]]:
        yield from buffered
        for row_number, cells in iterator:
            if len(cells) > max_cols:
                table.issues.append({"code": "TOO_MANY_COLUMNS", "row": row_number, "columns": len(cells)})
                cells = cells[:max_cols]
            yield row_number, cells

    for row_number, cells in source():
        if table.header_row_number is not None and row_number <= table.header_row_number:
            continue
        if all(c.value in (None, "") and c.formula is None for c in cells):
            continue
        if kind == FileKind.CSV and len(cells) != width and len(table.issues) < 1000:
            table.issues.append({"code": "INCONSISTENT_ROW_LENGTH", "row": row_number, "columns": len(cells), "expected": width})
        if counter["total"] >= max_rows:
            table.truncated = True
            counter["stop"] = True
            table.issues.append({"code": "ROW_LIMIT_REACHED", "row": row_number, "limit": max_rows})
            return
        record: dict[str, Any] = {}
        for col, cell in enumerate(cells):
            header = table.headers[col] if col < len(table.headers) else f"column_{col + 1}"
            value, vtype = cell_to_json(cell.value)
            if value is None and cell.formula is None:
                continue
            coord = f"{get_column_letter(col + 1)}{row_number}" if kind != FileKind.CSV else f"C{col + 1}"
            entry: dict[str, Any] = {"v": value, "t": vtype, "c": coord}
            if cell.formula:
                entry["f"] = cell.formula
                if cell.missing_cached:
                    entry["missing_cached"] = True
                    if len(table.issues) < 1000:
                        table.issues.append({"code": "FORMULA_CACHE_MISSING", "row": row_number, "cell": coord})
            record[header] = entry
        if record:
            counter["total"] += 1
            yield row_number, record


def rows_to_text(headers: list[str], rows: list[tuple[int, dict[str, Any]]], sheet: str | None) -> str:
    """Compact text rendering of rows for search indexing."""
    out = io.StringIO()
    for row_number, record in rows:
        parts = [f"{k}: {v['v']}" for k, v in record.items() if v.get("v")]
        prefix = f"[{sheet} row {row_number}]" if sheet else f"[row {row_number}]"
        out.write(prefix + " " + "; ".join(parts) + "\n")
    return out.getvalue()
